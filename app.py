import os
import json
import re
import base64
import fitz  # PyMuPDF
import docx  # python-docx
import openpyxl
import tempfile
import streamlit as st
from typing import Optional, List
from pathlib import Path

# --- 1. PAGE CONFIG (Must be first) ---
st.set_page_config(
    page_title="Gen AI PII Redaction", 
    layout="wide",
)

# --- 2. SAFE IMPORT BLOCK (Fixes your error) ---
try:
    import pypdf
except ImportError:
    import PyPDF2 as pypdf  # Fallback if the main one fails

# Required for structured output and LLM interaction
from pydantic import BaseModel, Field
from langchain_openai import ChatOpenAI
from langchain_core.messages import HumanMessage
from langchain_core.exceptions import OutputParserException


# --- Pydantic Schemas for Structured Output ---
class RedactionResult(BaseModel):
    """The result of a redaction task, containing the redacted text and a list of identified PII."""
    redacted_text: str = Field(description="The full document with all PII replaced by placeholders.")
    detected_pii: List[str] = Field(description="A list of all detected PII entities for logging purposes.")

# --- LLM and Agentic Pipeline Configuration ---

# UPDATED: Prioritize Streamlit Secrets, then Environment Variables.
if "OPENAI_API_KEY" in st.secrets:
    API_KEY = st.secrets["OPENAI_API_KEY"]
else:
    API_KEY = os.environ.get("OPENAI_API_KEY")

@st.cache_resource
def get_llm_chain():
    """Initializes and returns the structured LLM chain."""
    if not API_KEY:
        st.error("🚨 API Key missing! Please set OPENAI_API_KEY in Streamlit Secrets.")
        return None

    llm = ChatOpenAI(
        model="gpt-4o",
        temperature=0,
        # Ensure this base URL is correct for your provider (OpenRouter/OpenAI)
        openai_api_base="https://openrouter.ai/api/v1", 
        openai_api_key=API_KEY,
    )
    return llm.with_structured_output(RedactionResult)

redaction_llm = get_llm_chain()

# --- 1. Helper Function to Read File Content ---
def get_file_content(file_path: str, file_ext: str):
    """Extracts text content from various supported file types."""
    text_content = ""
    
    try:
        if file_ext == '.txt':
            with open(file_path, 'r', encoding='utf-8') as f:
                text_content = f.read()
        
        elif file_ext == '.docx':
            doc = docx.Document(file_path)
            text_content = "\n".join([para.text for para in doc.paragraphs])
        
        elif file_ext in ['.xlsx', '.xls']:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            for sheet in wb.worksheets:
                for row in sheet.iter_rows():
                    text_content += " ".join([str(cell.value) if cell.value is not None else "" for cell in row]) + "\n"
        
        elif file_ext == '.pdf':
            reader = pypdf.PdfReader(file_path)
            for page in reader.pages:
                text_content += page.extract_text() or ""
                
    except Exception as e:
        st.error(f"Error reading file {file_path}: {e}")
        return ""

    return text_content


# --- 2. Prompt Template ---
def get_redaction_prompt(file_path: str, file_ext: str, file_content: str) -> str:
    """Generates the multimodal or text-only prompt."""
    
    instruction = """
    You are a highly specialized and precise PII redaction agent. Your task is to analyze the document and return the full text with all PII replaced by specific, descriptive placeholders.

    Your goal is to preserve the original document's structure and formatting as much as possible while completely redacting sensitive information.

    Instructions for Redaction:
    1.  **Replace PII**: Find ALL parts of the Personal Name, including initials, prefixes (e.g., Mr., Dr.), and suffixes. Redact the entire name entity with the placeholder [NAME]. Apply similar complete redaction to all other PII entities.
    2.  **Use Placeholders**: Replace PII with the most appropriate placeholder from this list: [NAME], [DATE], [ID_NUMBER], [ADDRESS], [EMAIL], [PHONE_NUMBER], [FINANCIAL_INFO].
    3.  **Be Precise**: Only redact the PII itself. Do not redact surrounding text, punctuation, or formatting.
    4.  **Preserve Layout**: Maintain the original line breaks, spacing, and overall structure of the document.

    Return the full redacted text. Also, provide a list of the PII entities you detected for logging.
    """
    
    if file_ext in ['.txt', '.docx', '.xlsx', '.xls', '.pdf'] and file_content:
        # Text-enhanced prompt
        return f"{instruction}\n\nDocument text:\n---\n{file_content}\n---"
    else:
        # Multimodal prompt (for image files)
        return instruction


# --- 3. Main Processing Logic ---
def process_document_for_redaction(file_path: str) -> Optional[RedactionResult]:
    """Prepares the message and invokes the structured LLM chain."""
    if not redaction_llm: return None
    
    file_ext = os.path.splitext(file_path)[1].lower()
    
    is_image_or_pdf = file_ext in ['.pdf', '.jpg', '.jpeg', '.png', '.gif', '.tiff']
    is_text_based = file_ext in ['.txt', '.docx', '.xlsx', '.xls']
    
    if not (is_image_or_pdf or is_text_based):
        st.warning(f"Unsupported file type: {file_ext}")
        return None

    # 1. Get Text Content
    text_content = ""
    if is_text_based or file_ext == '.pdf':
        text_content = get_file_content(file_path, file_ext) or ""
    
    prompt = get_redaction_prompt(file_path, file_ext, text_content)
    content_parts = [{"type": "text", "text": prompt}]
    
    # 2. Add Image Data for Multimodal Files
    if is_image_or_pdf:
        try:
            img_bytes = None
            # For PDF, render the first page to PNG for vision processing
            if file_ext == '.pdf':
                DPI = 300
                doc = fitz.open(file_path)
                # Load first page only for vision context
                if len(doc) > 0:
                    page = doc.load_page(0)
                    pix = page.get_pixmap(dpi=DPI)
                    img_bytes = pix.tobytes("png")
                doc.close()
            else:
                with open(file_path, "rb") as image_file:
                    img_bytes = image_file.read()

            if img_bytes:
                base64_image = base64.b64encode(img_bytes).decode("utf-8")
                image_url = f"data:image/png;base64,{base64_image}"
                content_parts.append({"type": "image_url", "image_url": {"url": image_url}})

        except Exception as e:
            st.error(f"Error encoding {os.path.basename(file_path)} for vision model: {e}")
            return None
    
    # 3. Invoke LLM
    try:
        messages = [HumanMessage(content=content_parts)]
        result = redaction_llm.invoke(messages)
        return result
    except Exception as e:
        st.error(f"LLM processing failed: {e}")
        return None


# --- 4. Streamlit Application Block ---

def display_results_in_streamlit(file_name: str, result: RedactionResult):
    """Displays the redaction results in a clean, secure format."""
    
    st.header(f"Redaction Report for: **{file_name}**")
    
    col1, col2 = st.columns([1, 2])
    
    with col1:
        st.subheader("Detected PII Entities")
        if result.detected_pii:
            # --- SECURITY: Only show counts/types, not raw PII ---
            detected_types = set()
            count = 0
            for item in result.redacted_text.split():
                # Basic regex to catch [PLACEHOLDERS]
                match = re.search(r'\[([A-Z_]+)\]', item)
                if match:
                    detected_types.add(match.group(1))
            
            st.warning(f"Found {len(result.detected_pii)} PII items.")
            if detected_types:
                st.info(f"Redacted Types: {', '.join(sorted(list(detected_types)))}")
        else:
            st.success("No PII was detected.")
            
    with col2:
        st.subheader("Redaction Result")
        st.code(result.redacted_text, language='text')
        
    st.markdown("---")


def main_streamlit_app():
    # Page Config is now at the top of the file!
    
    # --- Logo and Title ---
    LOGO_PATH = Path("assets") / "an_logo.png"
    
    col_logo, col_title = st.columns([1, 4])
    with col_logo:
        if LOGO_PATH.exists():
            st.image(str(LOGO_PATH), width=150)
            
    with col_title:
        st.title("🤖 Automated PII Redaction Solution (Gen AI)")
        st.subheader("Redact Personal Identifiable Information from Documents")
        
    st.markdown("---")
    
    if not redaction_llm:
        st.error("LLM not initialized. Check API Key.")
        st.stop()
    
    # --- Mode Selection ---
    mode = st.radio(
        "Select Input Mode:",
        ("Upload Your Documents", "Run Demo (Uses Internal Samples)"),
        horizontal=True
    )
    
    files_to_process = []
    supported_extensions = ('.pdf', '.jpg', '.jpeg', '.png', '.gif', '.tiff', '.docx', '.xlsx', '.xls', '.txt')
    
    # 1. Upload Mode
    if mode == "Upload Your Documents":
        uploaded_files = st.file_uploader(
            "Upload single files or select a batch of files (PDF, Image, MS Office, TXT):",
            type=[ext.strip('.') for ext in supported_extensions],
            accept_multiple_files=True
        )
        
        if not uploaded_files:
            st.info("⬆️ Please upload documents to start processing.")
            return

        if st.button("Start Redaction", type="primary"):
            st.header("⚡ Processing Results")
            with tempfile.TemporaryDirectory() as temp_dir:
                for uploaded_file in uploaded_files:
                    file_name = uploaded_file.name
                    temp_file_path = os.path.join(temp_dir, file_name)
                    with open(temp_file_path, "wb") as f:
                        f.write(uploaded_file.getbuffer())
                    files_to_process.append((temp_file_path, file_name))
                    
                # Process loop
                for file_path, file_name in files_to_process:
                    with st.expander(f"**Analyzing {file_name}**", expanded=False):
                        with st.spinner(f"Redacting **{file_name}**..."):
                            result = process_document_for_redaction(file_path) 
                        
                        if result:
                            display_results_in_streamlit(file_name, result)
                        else:
                            st.error(f"❌ Failed to process **{file_name}**.")
                
                st.balloons()
                st.success("✅ All documents processed!")
                

    # 2. Demo Mode
    elif mode == "Run Demo (Uses Internal Samples)":
        DEMO_ROOT_FOLDER = "demo_documents"
        
        if st.button("Run Demo on Internal Samples", type="primary"):
            st.header("⚡ Processing Results")
            
            if os.path.isdir(DEMO_ROOT_FOLDER):
                for root, _, files in os.walk(DEMO_ROOT_FOLDER):
                    for f in files:
                        file_path = os.path.join(root, f)
                        if f.lower().endswith(supported_extensions) and os.path.isfile(file_path):
                            files_to_process.append((file_path, f))
            else:
                st.error(f"❌ Error: Demo folder '{DEMO_ROOT_FOLDER}' not found.")
                return

            if not files_to_process:
                st.error(f"No documents found in '{DEMO_ROOT_FOLDER}'.")
                return
            
            st.success(f"Found {len(files_to_process)} sample documents.")
            
            for file_path, file_name in files_to_process:
                with st.expander(f"**Analyzing {file_name}**", expanded=False):
                    with st.spinner(f"Redacting **{file_name}**..."):
                        result = process_document_for_redaction(file_path) 
                    
                    if result:
                        display_results_in_streamlit(file_name, result)
                    else:
                        st.error(f"❌ Failed to process **{file_name}**.")

            st.balloons()
            st.success("✅ All demo documents processed!")

if __name__ == "__main__":
    main_streamlit_app()
